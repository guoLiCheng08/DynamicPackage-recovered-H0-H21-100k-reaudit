from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import generate_principles_detailed_design_docx as principles
import generate_detailed_design_docx as dds
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/dynamicpackage_194_work/dynamicpackage_recovered_work')
OUT = ROOT / 'analysis' / 'documentation' / 'DynamicPackage_单星仿真_基础原理_逐式符号注释_最终版_H0-H21_10万步验证.docx'
CJK = dds.CJK
BLUE = dds.BLUE

# 每个键必须与公式图片生成调用的 key 一一对应。每个表仅解释该公式出现的符号。
SYMBOLS = {
    'p_continuous': [
        ('y(t)', '连续系统状态向量', '33×1 double；分量单位见状态表', 'DpState / y[0..32]'),
        ('t', '连续仿真时间', 's', 'integration_time / TimeAdd'),
        ('f', '状态导数函数', '状态单位/s', 'dp_differential_equation_33'),
        ('u(t)', '外部设备命令', '0x78 B 命令帧', 'DpDeviceControlCommand'),
        ('e(t)', '时间/环境/外力等外部量', '结构/向量', 'DpCoreEnvironmentOutputs'),
        ('Θ', '固定模型参数与配置', '矩阵、常量、限制', 'DpCoreDefaultModel 与设备配置'),
        ('ΦΔt', '步长 Δt 下的离散传播映射', '状态→状态', 'dp_rk4_integrator_33'),
    ],
    'p_output': [
        ('xdev,k', '第 k 步设备内部状态', '设备对象集合', 'RWheel/MTQ/Thruster/SADA'),
        ('u_k', '第 k 步外部命令', '0x78 B', 'DpDeviceControlCommand'),
        ('D(·)', '设备离散更新函数', '对象→对象', 'dp_update_device_control'),
        ('y_k', '第 k 步 33 维连续状态', '33×1 double', 'DpState'),
        ('ΦRK4(·)', 'RK4 核心传播', '状态→状态', 'dp_core_default_rk4_step_with_terms_and_environment'),
        ('e_k', '本步环境量', '时间/向量/结构', '太阳、地磁、外力与日历'),
        ('z_k', '传感器测量与有效标志', 'DeviceMeasure 对象', 'dp_update_device_measure_recovered'),
        ('ξ_k', '本步随机噪声样本序列', '标量序列', 'C rand()/ran_gaussian'),
        ('o_k', '对外 main/IPC 输出', '544 B / 3000 B', 'dp_update_main_out_pack / dp_send_dyn_tele'),
        ('H(·), G(·)', '传感器映射、输出封包映射', '对象/状态→对象/字节', 'sensors / telemetry / IPC'),
    ],
    'p_cross_dot': [
        ('a,b', '两个三维向量', '3×1；单位由语境决定', 'DpVector / DpVec3'),
        ('a·b', '点积，表示投影或夹角余弦分子', '两向量单位之积', 'vector_dot'),
        ('a×b', '叉积，方向垂直于 a、b', '两向量单位之积', 'vector3_cross'),
        ('a_x,a_y,a_z', '向量 a 的三个分量', '与 a 同单位', 'DpVector backing[0..2]'),
        ('b_x,b_y,b_z', '向量 b 的三个分量', '与 b 同单位', 'DpVector backing[0..2]'),
    ],
    'p_rotation': [
        ('A', '正交方向余弦矩阵', '3×3，无量纲', '姿态/坐标变换矩阵'),
        ('A^T', '矩阵 A 的转置；正交旋转的逆', '3×3，无量纲', 'matrix_trans'),
        ('I', '3×3 单位矩阵', '3×3，无量纲', '数学常量'),
        ('v_I', '惯性/源坐标系中的向量分量', '3×1；依变量而定', 'GCI 等源系向量'),
        ('v_B', '卫星本体系中的向量分量', '3×1；同 v_I', 'body-rate/field/force'),
        ('A_{B←I}', '从惯性系到本体系的旋转', '3×3，无量纲', '由四元数构造的姿态矩阵'),
    ],
    'p_quat': [
        ('q=[q0,q1,q2,q3]^T', '标量在前的单位姿态四元数', '4×1，无量纲', 'y[0..3] / DpQuat'),
        ('qdot', '四元数时间导数', '1/s', 'derivative[0..3]'),
        ('Xi(q)', '由 q 构造的 4×3 运动学矩阵', '4×3，无量纲', 'quat_xi'),
        ('ω', '卫星本体系角速度', '3×1，rad/s', 'y[4..6] / body_rate'),
        ('||q||', '四元数欧氏范数', '无量纲', 'dp_rk4_integrator_33 末尾'),
    ],
    'p_jd': [
        ('y,m,d', '公历年、月、日', '整数/日', 'Time 组件'),
        ('h,min,s', '时、分、秒', 'h、min、s', 'Time 组件与小数秒残量'),
        ('JD', '儒略日连续时间坐标', 'day', 'Calc_JD'),
        ('floor(·)', '向下取整', '无量纲', 'floor'),
    ],
    'p_gmst': [
        ('θ_GMST', '格林尼治平恒星时角', 'rad', 'Calc_theta_GMST'),
        ('T', '自 J2000 起的儒略世纪', '无量纲', '(JD-2451545)/36525'),
        ('t_UTC', '当日 UTC 秒数', 's', '3600h+60min+s'),
        ('fmod(·,86400)', '把恒星秒归一化到一天', 's', 'fmod'),
        ('π/43200', '恒星秒到弧度比例', 'rad/s', '动态时间常量'),
    ],
    'p_frames': [
        ('r_GCI', '地心惯性系位置向量', '3×1，m', 'y[7..9]'),
        ('r_ECEF', '地球固连系位置向量', '3×1，m', 'GCI2ECEF 输出'),
        ('R_Z(θ)', '绕 Z 轴被动旋转', '3×3，无量纲', 'axis_vector2matrix / time'),
        ('P(T)', '岁差旋转矩阵', '3×3，无量纲', 'precession matrix'),
        ('θ_GMST,T', 'GMST 角和儒略世纪', 'rad、无量纲', 'Calc_theta_GMST / Calc_JD'),
    ],
    'p_newton': [
        ('F', '万有引力', '3×1，N', '物理基础；代码以加速度计算'),
        ('m', '航天器质量', 'kg', 'SpacecraftMass'),
        ('μ', '地球标准引力参数 GM', 'm³/s²', 'dynamic_orbit.c 常量'),
        ('r_vec', '地心位置向量', '3×1，m', 'r_GCI/ECEF 中间量'),
        ('r=||r_vec||', '地心距离模长', 'm', 'vector_norm'),
        ('a0', '中心引力加速度', '3×1，m/s²', 'orbit_dynamic'),
        ('rdot,v', '位置导数/速度', '3×1，m/s', 'y[10..12]'),
    ],
    'p_energy': [
        ('ε', '比机械能', 'm²/s²', '轨道要素推导中间量'),
        ('v', '速度向量', '3×1，m/s', 'y[10..12]'),
        ('r', '地心距离模长', 'm', '||position||'),
        ('μ', '地球标准引力参数', 'm³/s²', 'dynamic_orbit.c'),
        ('a', '轨道半长轴', 'm', 'orbit_elements[0]'),
        ('n', '平均角速度/平均运动', 'rad/s', 'sqrt(mu/a³)'),
        ('P', '椭圆轨道周期', 's', '推导量；当前模型不单独存储'),
    ],
    'p_kepler': [
        ('M', '平均近点角', 'rad', '轨道要素/Elements2PosVel_M'),
        ('E', '偏近点角', 'rad', 'M2E 输出/迭代变量'),
        ('e', '轨道偏心率', '无量纲', 'orbit_elements[1]'),
        ('E_n,E_{n+1}', 'Newton 迭代前后偏近点角', 'rad', 'M2E 局部变量'),
    ],
    'p_harmonic': [
        ('a', '总轨道加速度', '3×1，m/s²', 'derivative[10..12]'),
        ('a0', '中心引力项', '3×1，m/s²', 'orbit_dynamic'),
        ('a_J2,J3,J4', '二/三/四阶重力带谐加速度', '3×1，m/s²', 'orbit_dynamic 展开多项式'),
        ('F_external', '惯性系外力', '3×1，N', 'F_I_external / 推力器'),
        ('m', '航天器质量', 'kg', 'SpacecraftMass'),
    ],
    'p_euler': [
        ('H', '本体系总角动量', '3×1，kg·m²/s', 'Htotal / dynamics'),
        ('dH/dt|_I', '惯性系中的角动量变化率', '3×1，N·m', '角动量定理'),
        ('τ_ext', '外部合力矩', '3×1，N·m', 'torque/dynamics'),
        ('J', '本体系刚体惯量矩阵', '3×3，kg·m²', 'DpCoreDefaultModel.inertia'),
        ('ω,ωdot', '本体系角速度与导数', 'rad/s，rad/s²', 'y[4..6] / derivative[4..6]'),
        ('H_add', '飞轮等附加角动量', '3×1，kg·m²/s', 'wheel/SADA/flex 项'),
    ],
    'p_wheel_torque': [
        ('τ_RW', '飞轮对卫星本体的反作用力矩', '3×1，N·m', 'minus_total_torque'),
        ('H_w', '轮组角动量', '3×1，kg·m²/s', 'WheelGroup.angular_momentum'),
        ('Hdot_w', '轮组角动量变化率', '3×1，N·m', 'WheelGroup.torque'),
    ],
    'p_mag_torque': [
        ('τ_MTQ', '磁力矩器产生的本体系力矩', '3×1，N·m', 'UpdateTorque/桥接项'),
        ('m', '群组磁偶极矩', '3×1，A·m²或模型单位', 'MTQ_Group.group_moment'),
        ('B', '本体系地磁向量', '3×1，模型磁场单位', 'prior_magnetic_body'),
    ],
    'p_gravity_gradient': [
        ('τ_gg', '重力梯度力矩', '3×1，N·m', 'dynamic_torque.c'),
        ('μ', '地球标准引力参数', 'm³/s²', 'dynamic_orbit.c'),
        ('r', '轨道半径模长', 'm', '||position_gci||'),
        ('rhat_B', '本体系中的径向单位向量', '3×1，无量纲', '姿态投影后的 position 单位向量'),
        ('J', '本体系惯量矩阵', '3×3，kg·m²', '核心模型 inertia'),
    ],
    'p_modal': [
        ('M', '模态广义质量/归一化质量矩阵', '10×10，配置相关', '理论参照；当前配置已吸收进恢复矩阵'),
        ('D', '模态阻尼矩阵', '10×10，配置相关', '对应 modal_d/MD'),
        ('K', '模态刚度矩阵', '10×10，配置相关', '对应 modal_a/MA'),
        ('η,ηdot,ηddot', '模态位移、速度、加速度', '10×1，模型模态单位', 'y[13..32]/derivative'),
        ('Q_rigid,Q_SADA', '刚体耦合与驱动的广义载荷', '10×1', 'flex SADA/rigid 项'),
    ],
    'p_effective_inertia': [
        ('J_eff', '耦合后有效惯量矩阵', '3×3，kg·m²', 'dp_flex_effective_inertia_3x3'),
        ('J', '基础刚体惯量矩阵', '3×3，kg·m²', 'core model inertia'),
        ('C', '刚柔耦合矩阵', '3×10，配置相关', 'core model coupling'),
        ('b', '刚体角加速度 RHS', '3×1，N·m', 'flex rigid RHS'),
        ('α', '本体系角加速度', '3×1，rad/s²', 'derivative[4..6]'),
    ],
    'p_modal_current': [
        ('ηddot', '10 阶模态加速度', '10×1，模型模态单位/s²', 'derivative[23..32]'),
        ('C,C^T', '刚柔耦合矩阵及转置', '3×10、10×3', 'dynamic_flex.c'),
        ('α', '刚体角加速度', '3×1，rad/s²', 'flex 解算结果'),
        ('MA,MD', '当前恢复模型的模态刚度/阻尼相关矩阵', '10×10，配置相关', 'modal_a/modal_d'),
        ('η,ηdot', '模态位移和模态速度', '10×1', 'y[13..32]'),
        ('optional terms', '已配置的 SADA/刚体额外模态项', '10×1', 'DpFlexSadaDrive/mapping'),
    ],
    'p_rk4_principle': [
        ('y', '待积分状态', '33×1 double', 'state[33]'),
        ('f', '状态导数函数', '状态单位/s', 'dp_differential_equation_33'),
        ('t', '当前积分时间', 's', 'integration_time'),
        ('h', '积分步长', 's', 'step_time'),
        ('k1..k4', '四个阶段的导数估计', '状态单位/s', 'RK4 阶段局部量'),
    ],
    'p_rk4_update': [
        ('y,y+', '积分前/后状态', '33×1 double', 'state 读写'),
        ('h', '积分步长', 's', 'step_time'),
        ('k1..k4', '四阶段导数', '状态单位/s', 'h_k1..h_k4 对应缓存'),
        ('q+', '积分后的四元数分量', '4×1，无量纲', 'state[0..3]'),
        ('||q+||', '积分后四元数范数', '无量纲', '归一化分母'),
    ],
    'p_actuator_wheel': [
        ('τ_w', '单轮扭矩', 'N·m', 'wheel.torque'),
        ('J_w', '单轮转子惯量', 'kg·m²', 'wheel.inertia'),
        ('α_w', '单轮角加速度', 'rad/s²', 'wheel.acceleration'),
        ('ω_w', '单轮角速度', 'rad/s', 'wheel.omega'),
        ('H_w', '单轮角动量', 'kg·m²/s', 'wheel.angular_momentum'),
        ('Δt', '设备离散更新步长', 's', 'device_control_step_time/step_time'),
    ],
    'p_actuator_thruster': [
        ('F_out', '推力器输出力', '3×1，N', 'Thruster.force_output'),
        ('k_F', '推力缩放系数', '无量纲', 'Thruster.force_scale'),
        ('F_in', '推力输入向量', '3×1，N或配置单位', 'Thruster.force_input'),
        ('τ_thr', '推力器力矩', '3×1，N·m', 'Thruster.torque_output'),
        ('l', '推力器力臂', '3×1，m', 'Thruster.lever_arm'),
        ('work_status', '推力器工作标志', 'uint32；1 为工作', 'Thruster.work_status'),
    ],
    'p_actuator_sada': [
        ('θ_cmd', 'SADA 命令角', 'rad或模型角度单位', 'SADA.command_angle'),
        ('θ_lim', 'SADA 命令角限制', '同 θ_cmd', 'SADA.command_limit'),
        ('ω_S', 'SADA 轴角速度', 'rad/s', 'SADA.angular_velocity'),
        ('Δω_S', '本步允许的角速度增量', 'rad/s', '局部计算'),
        ('a_max', '角加速度限制', 'rad/s²', 'SADA.acceleration_limit'),
        ('Δt', '设备步长', 's', 'drive_SADA_once 参数'),
        ('θ+', '更新后当前角', 'rad或模型单位', 'SADA.current_angle'),
    ],
    'p_measurement': [
        ('z', '传感器测量', '向量/四元数/标志', 'DeviceMeasure'),
        ('h(x)', '从真实状态到理想测量的非线性映射', '状态→测量', 'Gyro/Mag/DSS/STS/GPS 更新'),
        ('x', '真实状态与环境', '33维状态及环境', 'DpState / environment'),
        ('b', '确定性测量偏置', '同 z', 'gaussian_mean 等配置'),
        ('v', '随机测量噪声', '同 z', 'ran_gaussian2 输出'),
        ('R', '噪声协方差的理论记号', '测量维×测量维', '当前实现以独立 sigma 分量表达，不显式存储 R'),
    ],
    'p_gaussian': [
        ('x,y', '拒绝采样得到的均匀随机中间量', '无量纲，范围 (-1,1)', 'rand() 归一化结果'),
        ('s', '半径平方 x²+y²', '无量纲，(0,1]', '拒绝条件'),
        ('σ', '高斯标准差', '被测量的单位', 'gaussian_sigma / 函数参数'),
        ('μ', '高斯均值/偏置', '被测量的单位', 'gaussian_mean / 函数参数'),
        ('n', '零均值噪声样本', '同 σ', 'ran_gaussian'),
        ('z', '带均值样本', '同 μ、σ', 'ran_gaussian2'),
    ],
    'p_dss': [
        ('u', 'DSS 安装坐标中的太阳方向投影', '3×1，无量纲/方向量', 'DSS.measure 中间量'),
        ('P_DSS', 'DSS 投影/安装矩阵', '3×3，无量纲', 'DSS.projection_matrix'),
        ('s_B', '本体系太阳方向/向量', '3×1，方向/模型距离', 'environment.sun_body'),
        ('u_x,u_y,u_z', '投影向量三个分量', '同 u', '局部/measure'),
        ('alpha_x,alpha_y', 'DSS 横向角输出', 'rad，遥测时转换为度', 'DSS.angle_x/angle_y'),
    ],
    'p_visibility': [
        ('alpha_limb', '从卫星看地球圆盘的边缘角', 'rad', 'isEarthShadow'),
        ('R_e', '地球参考半径', 'm', 'orbit/environment 常量'),
        ('r,r_GCI', '位置模长与 GCI 位置向量', 'm；3×1 m', 'state.position_gci'),
        ('s_GCI', '太阳 GCI 向量', '3×1，m/方向', 'environment.sun_gci'),
        ('angle(·,·)', '两向量夹角', 'rad', 'vector2angle'),
        ('shadow', '地影布尔判定', 'int/布尔', 'isEarthShadow 返回值'),
    ],
    'p_serialization': [
        ('bytes[off:off+n]', '目标帧中的一段字节', 'n B', 'main/IPC raw buffer'),
        ('off,n', '字段偏移与字段字节数', 'B', 'telemetry/IP C layout 常量'),
        ('encode_endian,type', '指定端序和标量类型的编码', 'value→bytes', 'var2buff/Put_*_In_Buff'),
        ('value_double,value_float', '双精度源值与 float32 协议值', 'double、float', 'sendDynTele setter'),
    ],
    'p_difference': [
        ('gold_{k,b}', '原 ELF 在第 k 步、第 b 块的金标字节', '字节数组', 'gold_*.bin'),
        ('recovered_{k,b}', '恢复端在同一步/块的字节', '字节数组', '比较器当前输出'),
        ('memcmp(·)', '逐字节比较函数', 'int；0 表示相同', 'C memcmp'),
        ('Δ_{k,b}', '该步该块的差异结果', 'int', '比较器诊断'),
        ('PASS', '所有被观察 k,b 均无差异', '布尔', 'selftest gate'),
    ],
    'dds_single_step': [
        ('xdev,k+1', '设备更新后的对象状态', '设备对象集合', 'dp_update_device_control 后'),
        ('D', '设备更新算子', '对象→对象', 'dp_update_device_control'),
        ('y_k,y_k+1', '前/后 33维连续状态', '33×1 double', 'DpState'),
        ('Φ_RK4', '核心 RK4 状态传播', '状态→状态', 'core default RK4'),
        ('t_k,Δt', '积分时间与步长', 's', 'integration_time / step_time'),
        ('e_k,e_k+1', '环境量', '结构/向量', 'core environment'),
        ('z_k+1,H', '传感器输出及其映射', '对象/函数', 'dp_update_device_measure_recovered'),
        ('o_k+1,G', '遥测/IPC 输出及封包映射', '字节帧/函数', 'telemetry / IPC'),
        ('Θ,ξ_k', '配置与随机样本', '参数集合/标量序列', 'model/device config；rand'),
    ],
    'dds_vector': [
        ('a,b', '两个三维向量', '3×1；依语境', 'DpVector'),
        ('a×b', '叉积', '3×1；单位相乘', 'vector3_cross'),
        ('a^T b', '点积', '标量；单位相乘', 'vector_dot'),
        ('||a||,||b||', '向量欧氏范数', '与向量同单位', 'vector_norm'),
        ('θ', '向量夹角', 'rad', 'vector2angle'),
    ],
    'dds_quat': [
        ('q0..q3', '标量在前的四元数分量', '无量纲', 'state[0..3]'),
        ('Xi(q)', '4×3 四元数运动学矩阵', '4×3，无量纲', 'quat_xi'),
        ('qdot', '四元数导数', '1/s', 'differential_equation output'),
        ('ω', '本体系角速度', '3×1，rad/s', 'state[4..6]'),
    ],
    'dds_jd_gmst': [
        ('y,m,d,h,min,s', '公历日期与时间分量', '年、月、日、时、分、秒', 'Time 状态'),
        ('JD', '儒略日', 'day', 'Calc_JD'),
        ('floor(·)', '向下取整', '无量纲', 'floor'),
        ('T', '儒略世纪', '无量纲', '(JD-2451545)/36525'),
    ],
    'dds_gci_ecef': [
        ('r_GCI,r_ECEF', '惯性/地固位置向量', '各 3×1 m', 'state 与坐标转换输出'),
        ('R_Z', 'Z轴旋转矩阵', '3×3 无量纲', 'time rotations'),
        ('θ_GMST', '恒星时角', 'rad', 'Calc_theta_GMST'),
        ('P(T)', '岁差矩阵', '3×3 无量纲', 'precession'),
        ('T', '儒略世纪', '无量纲', 'JD 转换'),
    ],
    'dds_kepler': [
        ('E_k,E_k+1', 'Newton 迭代前/后偏近点角', 'rad', 'M2E'),
        ('M', '平均近点角', 'rad', '轨道要素'),
        ('e', '偏心率', '无量纲', 'orbit element'),
        ('ΔE', '迭代改变量', 'rad', '收敛判定 <1e-10'),
    ],
    'dds_orbit_acc': [
        ('a,a0,a_J2,a_J3,a_J4', '总/中心/各阶带谐加速度', '3×1 m/s²', 'orbit_dynamic'),
        ('F_external', '惯性外力', '3×1 N', 'F_I_external'),
        ('m', '质量', 'kg', 'SpacecraftMass'),
        ('μ', '引力参数', 'm³/s²', 'dynamic_orbit.c'),
        ('r_vec,||r||', '位置向量及模长', 'm', 'position state'),
    ],
    'dds_sun': [
        ('M', '太阳模型平均近点角相位', 'rad', 'dp_sun_vector'),
        ('T', '儒略世纪', '无量纲', 'JD 时间'),
        ('frac(·)', '取小数部分的周期相位函数', '无量纲', '数学辅助'),
        ('r_sun', '日地距离近似', 'm', 'sun vector magnitude'),
    ],
    'dds_total_h': [
        ('H_total', '系统总角动量', '3×1 kg·m²/s', 'dynamics'),
        ('J', '基础惯量矩阵', '3×3 kg·m²', 'core model inertia'),
        ('ω', '本体系角速度', '3×1 rad/s', 'state[4..6]'),
        ('H_w', '轮组角动量', '3×1 kg·m²/s', 'WheelGroup'),
        ('H_SADA,cmd', 'SADA 命令映射动量', '3×1 kg·m²/s', 'sada command momentum map'),
        ('C,η', '耦合矩阵和模态位移', '3×10；10×1', 'core coupling / state[13..22]'),
        ('J_eff', '有效惯量', '3×3 kg·m²', 'flex effective inertia'),
    ],
    'dds_flex_rhs': [
        ('α', '刚体角加速度', '3×1 rad/s²', 'dynamics/flex'),
        ('J_eff,b', '有效惯量与 RHS', '3×3 kg·m²；3×1 N·m', 'flex solve'),
        ('C,C^T', '刚柔耦合矩阵及转置', '3×10；10×3', 'core coupling'),
        ('M4,u_rigid', '可选刚体耦合矩阵和输入', '配置相关；3×1', 'flex optional term'),
        ('C_S,P_S', 'SADA 模态预映射与载荷', '10×3；3×1', 'SADA flex maps'),
        ('MA,MD', '模态刚度/阻尼相关矩阵', '10×10', 'modal_a/modal_d'),
        ('η,ηdot,ηddot', '模态位移/速度/加速度', '10×1', 'state/derivative'),
    ],
    'dds_rk4': [
        ('k1..k4', '四个阶段导数', '状态单位/s', 'RK4 局部阶段'),
        ('f', '33 维导数函数', '状态单位/s', 'dp_differential_equation_33'),
        ('t,h', '积分时间与步长', 's', 'integration_time/step_time'),
        ('y,y+', '积分前后状态', '33×1 double', 'state'),
    ],
    'dds_wheel': [
        ('α_i^cand', '第 i 轮候选角加速度', 'rad/s²', 'dp_set_wheel_acc'),
        ('τ_i^cmd,τ_i', '轮扭矩命令/实际扭矩', 'N·m', 'command/wheel.torque'),
        ('J_i', '第 i 轮惯量', 'kg·m²', 'wheel.inertia'),
        ('χ_i', '噪声开关', '0/1', 'add_gaussian_noise_flag'),
        ('N(0,σ)', '高斯样本', 'rad/s²', 'ran_gaussian'),
        ('ω_i,H_i,Δt', '轮速、角动量、步长', 'rad/s；kg·m²/s；s', 'wheel state / device step'),
    ],
    'dds_mtq': [
        ('m_i^cmd,m_i^act', '第 i 个 MTQ 命令/饱和后磁矩', 'A·m²或模型单位', 'command/MTQ.actual_moment'),
        ('m_i,max', '第 i 通道磁矩上限', '同 m', 'MTQ.moment_limit'),
        ('sat(·)', '对称饱和函数', '同输入', 'dp_set_mtq_moment'),
        ('M_3×6', '六通道到三维群组的安装矩阵', '3×6，无量纲', 'MTQ_Group.mapping_3x6'),
        ('m_group,B,τ_MTQ', '群组磁矩、本体系磁场、产生力矩', '3×1；3×1；3×1 N·m', 'MTQ_Group/prior B/torque'),
    ],
    'dds_thruster': [
        ('work_status', '推力器开关', 'uint32，1为工作', 'Thruster.work_status'),
        ('F_out,F_in', '输出/输入推力向量', '3×1 N或模型单位', 'Thruster.force_output/force_input'),
        ('k_F', '推力缩放', '无量纲', 'Thruster.force_scale'),
        ('τ_out', '推力器输出力矩', '3×1 N·m', 'Thruster.torque_output'),
        ('l', '力臂', '3×1 m', 'Thruster.lever_arm'),
    ],
    'dds_sada': [
        ('θ_cmd,θ_lim', '命令角与限位', '角度单位', 'SADA.command_angle/limit'),
        ('sat(·)', '限幅函数', '同输入', 'drive_SADA_once'),
        ('ω,Δω,ω+', '当前角速度、增量、更新速度', 'rad/s', 'SADA.angular_velocity'),
        ('a_lim,a', '加速度限制与实际加速度', 'rad/s²', 'SADA.acceleration_limit/angular_acceleration'),
        ('θ+', '更新后的角度', '角度单位', 'SADA.current_angle'),
        ('Δt', '设备离散步长', 's', 'drive_SADA_once'),
    ],
    'dds_gaussian': [
        ('x,y,s', '拒绝采样中间变量及半径平方', '无量纲', 'ran_gaussian'),
        ('σ', '标准差', '被测量单位', '函数参数'),
        ('n', '零均值高斯噪声', '同 σ', '函数返回中间量'),
        ('μ,z', '均值/带均值随机样本', '被测量单位', 'ran_gaussian2'),
    ],
    'dds_gyro': [
        ('z', '陀螺测量向量', '3×1 rad/s', 'Gyro.measure'),
        ('P', '陀螺安装投影矩阵', '3×3 无量纲', 'Gyro.projection_matrix'),
        ('ω', '本体系真实角速度', '3×1 rad/s', 'state.body_rate'),
        ('χ_noise', '噪声使能开关', '0/1', 'Gyro.gaussian_noise_flag'),
        ('μ_i,σ_i', '第 i 分量噪声均值/标准差', 'rad/s', 'Gyro gaussian arrays'),
    ],
    'dds_mag': [
        ('z_mag,j', '第 j 路磁强计测量', '3×1 磁场单位', 'MagMeter[j].measure'),
        ('P_mag,j', '第 j 路磁强计投影矩阵', '3×3 无量纲', 'MagMeter[j].projection_matrix'),
        ('B_body', '本体系真实地磁', '3×1 磁场单位', 'environment.magnetic_body'),
        ('χ_j,μ_j,σ_j', '对应噪声开关、均值、标准差', '0/1、磁场单位', 'MagMeter 配置'),
        ('j', '磁强计索引', '{0,1}', 'MagMeter[2]'),
    ],
    'dds_dss': [
        ('u', 'DSS 安装系太阳投影', '3×1 方向量', 'DSS 中间测量'),
        ('P_DSS', 'DSS 投影矩阵', '3×3 无量纲', 'DSS.projection_matrix'),
        ('s_B', '本体系太阳向量', '3×1 方向/模型距离', 'sun_body'),
        ('u_z', '前向分量', '同 u；须>0', '投影结果'),
        ('u_tilde_x,u_tilde_y', '以 z 归一化的横向比值', '无量纲', 'DSS.ratio_x/y'),
        ('α_x,α_y', '太阳横向角', 'rad（遥测转换为度）', 'DSS.angle_x/y'),
    ],
    'dds_shadow': [
        ('alpha_limb', '地球边缘角', 'rad', 'isEarthShadow'),
        ('R_e', '地球半径', 'm', 'orbit constant'),
        ('r,r_GCI', '位置模长/GCI位置', 'm；3×1 m', 'state.position_gci'),
        ('s_GCI', '太阳 GCI 向量', '3×1', 'environment.sun_gci'),
        ('shadow', '阴影判定', 'int/布尔', 'isEarthShadow'),
        ('θ_DSS,z_DSS', 'DSS 有效性夹角/测量向量', 'rad；3×1', 'UpdateSunSensorValidFlag/DSS.measure'),
        ('(0,0,1)^T', '固定 DSS 参考轴', '3×1 无量纲', 'H11 原 ELF 调用点证据'),
    ],
    'dds_sts': [
        ('alpha_E', '地球视角半角', 'rad', 'isEarthOut/isStarTrackerValid'),
        ('R_e,r', '地球半径与位置模长', 'm', 'orbit/state position'),
        ('q_meas', '星敏感器测量四元数', '4×1 无量纲', 'STS.measure_quat'),
        ('q_error', '星敏误差四元数', '4×1 无量纲', 'STS.error_quat'),
        ('q_truth', '真实姿态四元数', '4×1 无量纲', 'state attitude'),
        ('⊗', '四元数乘法', '4×1×4×1→4×1', 'quat_cross ABI 适配器'),
    ],
}

USED = []
ORIGINAL_ADD_FORMULA = dds.base.add_formula_image


def add_symbol_table(doc, key):
    rows = SYMBOLS.get(key)
    if rows is None:
        raise RuntimeError(f'缺少公式 {key} 的就地符号表')
    # 公式标题、图片、原 caption 后紧接此表；不要求读者回查全局符号表。
    dds.base.add_table(doc, ['符号（本式）', '含义', '维度/单位/坐标系', '当前代码映射'], rows,
                       widths=[3.0, 5.1, 4.3, 4.8], font_size=7.55)


def annotated_formula(doc, key, latex, caption, width_cm=14.6):
    ORIGINAL_ADD_FORMULA(doc, key, latex, caption, width_cm)
    add_symbol_table(doc, key)
    USED.append(key)


def set_font(run, size=10, bold=False, color=None, name=CJK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_cover(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(54); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('DynamicPackage'); set_font(r, name='Noto Serif CJK SC', size=28, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('单星仿真基础原理、详细设计与逐式符号注释'); set_font(r, name='Noto Serif CJK SC', size=19.5, bold=True, color='0F6B78')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('规则：每一条公式之后立即解释本式出现的全部符号｜H0–H21（含 100,000 步验证）'); set_font(r, size=10.5, color='666666')
    doc.add_paragraph()
    dds.base.add_callout(doc, '本版阅读规则', '从本页后的所有公式图片开始，公式编号/说明文字下方紧接“符号（本式）—含义—维度/单位/坐标系—当前代码映射”表。表内只解释该公式实际使用的符号；读者无需翻回统一符号表。相同符号在不同公式中会按当前语境重复解释，避免歧义。', fill='E2F0D9')
    dds.base.add_table(doc, ['核对项', '本版处理'], [
        ('公式数量', '对基础原理与详细设计中所有公式调用逐键注释；生成时自动检查没有遗漏 key。'),
        ('符号范围', '包括数学变量、物理量、离散下标、矩阵、布尔标志、饱和函数、编码偏移和协议字节。'),
        ('代码追溯', '每个符号表均标出具体状态槽位、对象字段、函数或配置矩阵。'),
        ('阅读边界', '数学原理与当前实现映射并列；常量、分支、浮点顺序最终仍以工程源码和原 ELF gold 为准。'),
    ], widths=[3.5, 13.7], font_size=8.8)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(26)
    r = p.add_run('仅限离线研究与仿真；不得用于飞行、控制、实体执行机构或安全关键用途。'); set_font(r, size=9.5, bold=True, color='9C0006')
    doc.add_page_break()


def build_doc():
    # 所有旧章节调用的 dds.base.add_formula_image 将被本地包装器替代。
    dds.base.add_formula_image = annotated_formula
    try:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(1.45); section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.45); section.right_margin = Cm(1.45)
        dds.base.create_styles(doc)
        footer = section.footer.paragraphs[0]; footer.text = ''
        dds.base.add_page_number(footer)
        header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = header.add_run('DynamicPackage 逐式符号注释详细设计  |  H0–H21（含 100,000 步）'); set_font(r, size=8, color='7F7F7F')

        add_cover(doc)
        principles.add_principles(doc)
        dds.add_reader_contract(doc)
        dds.add_architecture(doc)
        dds.add_math_module(doc)
        dds.add_time_coordinate_module(doc)
        dds.add_orbit_environment_modules(doc)
        dds.add_dynamics_modules(doc)
        dds.add_device_modules(doc)
        dds.add_sensor_modules(doc)
        dds.add_output_modules(doc)
        dds.add_verification(doc)
        dds.add_appendices(doc)

        used = set(USED)
        defined = set(SYMBOLS)
        missing = defined - used
        unexpected = used - defined
        if missing or unexpected:
            raise RuntimeError(f'公式符号注释一致性失败：missing={sorted(missing)}, unexpected={sorted(unexpected)}')
        doc.core_properties.title = 'DynamicPackage 单星仿真基础原理、详细设计与逐式符号注释（H0–H21，含 100,000 步验证）'
        doc.core_properties.subject = '每条公式就地解释符号、单位、维度、坐标系与代码映射'
        doc.core_properties.author = 'Manus AI'
        doc.core_properties.comments = '仅限离线研究与仿真。'
        doc.save(OUT)
        print(f'annotated_formula_calls={len(USED)}')
        print(f'unique_formula_keys={len(used)}')
        print(OUT)
    finally:
        dds.base.add_formula_image = ORIGINAL_ADD_FORMULA


if __name__ == '__main__':
    build_doc()
