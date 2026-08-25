from pathlib import Path
import re
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import generate_formula_annotated_docx as src
import generate_formula_annotated_toc_docx as toc
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/dynamicpackage_194_work/dynamicpackage_recovered_work')
OUT = ROOT / 'analysis' / 'documentation' / 'DynamicPackage_单星仿真_基础原理_逐式符号注释_WPS可编辑线性公式_H0-H20.docx'
WPS_FONT = 'Microsoft YaHei'
USED = []

ASCII_REPLACEMENTS = {
    'α': 'alpha', 'β': 'beta', 'γ': 'gamma', 'δ': 'delta', 'Δ': 'Delta',
    'ε': 'epsilon', 'η': 'eta', 'θ': 'theta', 'Θ': 'Theta', 'λ': 'lambda',
    'μ': 'mu', 'ν': 'nu', 'ξ': 'xi', 'Ξ': 'Xi', 'π': 'pi', 'ρ': 'rho',
    'σ': 'sigma', 'Σ': 'Sigma', 'τ': 'tau', 'φ': 'phi', 'Φ': 'Phi',
    'χ': 'chi', 'ψ': 'psi', 'Ψ': 'Psi', 'ω': 'omega', 'Ω': 'Omega',
    '⊗': 'otimes', '×': 'x', '→': '->', '←': '<-', '↔': '<->',
    '≤': '<=', '≥': '>=', '≠': '!=', '∈': 'in', '∥': '||', '·': '*',
    'ᵀ': '_T', '⃗': '_vec', '…': '...', '–': '-', '—': '-',
}

LATEX_REPLACEMENTS = [
    ('\\left', ''), ('\\right', ''), ('\\mathbf', ''), ('\\boldsymbol', ''),
    ('\\mathrm', ''), ('\\text', ''), ('\\quad', '  '), ('\\;', ' '),
    ('\\,', ' '), ('\\!', ''), ('\\cdot', '*'), ('\\times', ' x '),
    ('\\to', ' -> '), ('\\mapsto', ' -> '), ('\\leq', '<='), ('\\geq', '>='),
    ('\\neq', '!='), ('\\approx', '~'), ('\\in', ' in '), ('\\partial', 'd'),
    ('\\sum', 'sum'), ('\\sqrt', 'sqrt'), ('\\arccos', 'acos'),
    ('\\sin', 'sin'), ('\\cos', 'cos'), ('\\tan', 'tan'), ('\\log', 'log'),
    ('\\exp', 'exp'), ('\\hat', ''), ('\\bar', ''), ('\\dot', 'dot'),
    ('\\ddot', 'ddot'), ('\\frac', 'frac'), ('\\lVert', '||'), ('\\rVert', '||'),
    ('\\Vert', '||'), ('\\|', '||'), ('\\begin', ''), ('\\end', ''),
]


def set_wps_font(run, size=9.0, bold=False, color=None):
    run.font.name = WPS_FONT
    run._element.rPr.rFonts.set(qn('w:ascii'), WPS_FONT)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), WPS_FONT)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), WPS_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color


def ascii_safe(value):
    text = str(value)
    for old, new in ASCII_REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def linearize_latex(latex):
    text = latex.strip()
    for old, new in LATEX_REPLACEMENTS:
        text = text.replace(old, new)
    greek = {
        '\\alpha':'alpha','\\beta':'beta','\\gamma':'gamma','\\delta':'delta',
        '\\epsilon':'epsilon','\\eta':'eta','\\theta':'theta','\\lambda':'lambda',
        '\\mu':'mu','\\nu':'nu','\\xi':'xi','\\pi':'pi','\\rho':'rho',
        '\\sigma':'sigma','\\tau':'tau','\\phi':'phi','\\chi':'chi','\\psi':'psi','\\omega':'omega',
        '\\Theta':'Theta','\\Phi':'Phi','\\Xi':'Xi','\\Psi':'Psi','\\Omega':'Omega',
    }
    for old, new in greek.items():
        text = text.replace(old, new)
    # 将常见 TeX 下标/上标转换为可编辑的线性写法。
    text = re.sub(r'_\{([^{}]+)\}', r'_(\1)', text)
    text = re.sub(r'\^\{([^{}]+)\}', r'^(\1)', text)
    text = re.sub(r'_([A-Za-z0-9+\-])', r'_\1', text)
    text = re.sub(r'\^([A-Za-z0-9+\-])', r'^\1', text)
    text = re.sub(r'frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1)/(\2)', text)
    text = text.replace('{', '(').replace('}', ')')
    text = text.replace('[(', '[').replace(')]', ']')
    text = re.sub(r'\\([A-Za-z]+)', r'\1', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return ascii_safe(text)


def add_wps_symbol_table(doc, key):
    rows = src.SYMBOLS.get(key)
    if rows is None:
        raise RuntimeError(f'缺少公式 {key} 的符号表')
    converted = [tuple(ascii_safe(cell) for cell in row) for row in rows]
    src.dds.base.add_table(
        doc,
        ['Symbol (ASCII)', 'Meaning', 'Dimension / unit / frame', 'Current code mapping'],
        converted, widths=[3.1, 5.0, 4.35, 4.75], font_size=7.5
    )


def add_linear_formula(doc, key, latex, caption, width_cm=14.6):
    """用普通可编辑文本替代 PNG/OMML 公式。"""
    text = linearize_latex(latex)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    src.dds.base.set_cell_shading(cell, 'F2F2F2')
    src.dds.base.set_cell_margins(cell, top=95, start=120, bottom=95, end=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_wps_font(r, size=8.6, bold=False)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(2)
    r = cap.add_run(ascii_safe(caption) + '  [Editable linear formula; no OMML / no image]')
    set_wps_font(r, size=7.4, color='666666')
    add_wps_symbol_table(doc, key)
    USED.append(key)


def add_wps_code(doc, title, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.cell(0, 0)
    src.dds.base.set_cell_shading(cell, 'F3F6F8')
    src.dds.base.set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(ascii_safe(title))
    set_wps_font(r, size=9.2, bold=True)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(ascii_safe(line))
        # 不使用 Courier/DejaVu 等等宽字体，避免 WPS 映射成异常字形。
        set_wps_font(r, size=8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def build_doc():
    original_formula = src.dds.base.add_formula_image
    original_code = src.dds.add_code
    src.USED.clear()
    try:
        src.dds.base.add_formula_image = add_linear_formula
        src.dds.add_code = add_wps_code
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(1.45); section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.45); section.right_margin = Cm(1.45)
        src.dds.base.create_styles(doc)
        footer = section.footer.paragraphs[0]; footer.text = ''
        src.dds.base.add_page_number(footer)
        header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h = header.add_run('DynamicPackage WPS editable linear formulas | H0-H20')
        set_wps_font(h, size=8.0)

        # 自定义封面，明确该版不使用图片公式、OMML 或等宽代码字体。
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(54); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('DynamicPackage'); set_wps_font(r, size=28, bold=True)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Single-satellite principles, detailed design, and editable linear formulas'); set_wps_font(r, size=16, bold=True)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('WPS compatibility edition: no OMML, no equation images, no Courier/DejaVu code font'); set_wps_font(r, size=9.5)
        src.dds.base.add_callout(doc, 'WPS compatibility rule', 'All formulas in this edition are normal editable text in ASCII linear form. Select a formula line and copy/edit it directly. If native WPS professional formatting is required, paste the linear text into WPS Insert > Equation on the target machine; the current document intentionally contains no OMML object.', fill='E2F0D9')
        doc.add_page_break()
        toc.add_auto_toc(doc)

        src.principles.add_principles(doc)
        src.dds.add_reader_contract(doc)
        src.dds.add_architecture(doc)
        src.dds.add_math_module(doc)
        src.dds.add_time_coordinate_module(doc)
        src.dds.add_orbit_environment_modules(doc)
        src.dds.add_dynamics_modules(doc)
        src.dds.add_device_modules(doc)
        src.dds.add_sensor_modules(doc)
        src.dds.add_output_modules(doc)
        src.dds.add_verification(doc)
        src.dds.add_appendices(doc)

        used = set(USED)
        defined = set(src.SYMBOLS)
        missing = defined - used
        unexpected = used - defined
        if missing or unexpected:
            raise RuntimeError(f'Formula annotation mismatch: missing={sorted(missing)}, unexpected={sorted(unexpected)}')
        toc.set_update_fields_on_open(doc)
        doc.core_properties.title = 'DynamicPackage WPS editable linear formulas (H0-H20)'
        doc.core_properties.subject = 'No OMML; no equation images; editable ASCII linear formulas'
        doc.core_properties.author = 'Manus AI'
        doc.save(OUT)
        print(f'editable_linear_formula_calls={len(USED)}')
        print(OUT)
    finally:
        src.dds.base.add_formula_image = original_formula
        src.dds.add_code = original_code


if __name__ == '__main__':
    build_doc()
