from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import generate_formula_annotated_docx as src
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/dynamicpackage_194_work/dynamicpackage_recovered_work')
OUT = ROOT / 'analysis' / 'documentation' / 'DynamicPackage_单星仿真_基础原理_逐式符号注释_最终版_Word自动更新目录_H0-H21_10万步验证.docx'


def set_update_fields_on_open(doc):
    """请求 Word/WPS 在打开文档时更新域，包括 TOC 页码。"""
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)


def add_field(run, field_type, text=None):
    node = OxmlElement('w:fldChar') if field_type else OxmlElement('w:instrText')
    if field_type:
        node.set(qn('w:fldCharType'), field_type)
    else:
        node.set(qn('xml:space'), 'preserve')
        node.text = text
    run._r.append(node)


def add_auto_toc(doc):
    """插入标准 Word TOC 域；目录项按标题级别 1–3 自动生成。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('目录')
    src.set_font(r, name='Noto Serif CJK SC', size=20, bold=True, color='0F6B78')

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(8)
    rn = note.add_run('本目录为 Word/WPS 自动目录：打开文档时将自动更新；若未更新，请在目录上右键选择“更新域/更新目录”。')
    src.set_font(rn, size=8.5, color='666666')

    toc = doc.add_paragraph()
    toc.paragraph_format.space_after = Pt(6)
    r = toc.add_run()
    add_field(r, 'begin')
    add_field(r, None, ' TOC \\o "1-3" \\h \\z \\u ')
    add_field(r, 'separate')
    fallback = toc.add_run('目录正在更新。若未自动显示，请在 Word/WPS 中右键目录并选择“更新整个目录”。')
    src.set_font(fallback, size=10, color='666666')
    r_end = toc.add_run()
    add_field(r_end, 'end')
    doc.add_page_break()


def build_doc():
    src.dds.base.add_formula_image = src.annotated_formula
    src.USED.clear()
    try:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Cm(1.45); section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.45); section.right_margin = Cm(1.45)
        src.dds.base.create_styles(doc)
        footer = section.footer.paragraphs[0]; footer.text = ''
        src.dds.base.add_page_number(footer)
        header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rh = header.add_run('DynamicPackage 自动目录与逐式符号注释  |  H0–H21（含 100,000 步）')
        src.set_font(rh, size=8, color='7F7F7F')

        src.add_cover(doc)
        add_auto_toc(doc)
        src.principles.add_principles(doc)
        src.dds.add_reader_contract(doc)
        src.dds.add_architecture(doc)
        src.dds.add_math_module(doc)
        src.dds.add_time_coordinate_module(doc)
        src.dds.add_orbit_environment_modules(doc)
        src.dds.add_dynamics_modules(doc)
        src.dds.add_device_modules(doc)
        src.dds.add_sensor_modules(doc)
        src.dds.add_output_modules(doc)
        src.dds.add_verification(doc)
        src.dds.add_appendices(doc)

        used = set(src.USED)
        defined = set(src.SYMBOLS)
        missing = defined - used
        unexpected = used - defined
        if missing or unexpected:
            raise RuntimeError(f'公式符号注释一致性失败：missing={sorted(missing)}, unexpected={sorted(unexpected)}')

        set_update_fields_on_open(doc)
        doc.core_properties.title = 'DynamicPackage 单星仿真基础原理、详细设计、逐式符号注释与自动目录（H0–H21，含 100,000 步验证）'
        doc.core_properties.subject = 'Word/WPS 可自动更新目录；逐式符号注释'
        doc.core_properties.author = 'Manus AI'
        doc.core_properties.comments = '仅限离线研究与仿真。'
        doc.save(OUT)
        print(f'annotated_formula_calls={len(src.USED)}')
        print(f'unique_formula_keys={len(used)}')
        print(OUT)
    finally:
        src.dds.base.add_formula_image = src.ORIGINAL_ADD_FORMULA


if __name__ == '__main__':
    build_doc()
