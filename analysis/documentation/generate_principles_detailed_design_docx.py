from pathlib import Path
import sys

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import generate_detailed_design_docx as dds
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path('/home/ubuntu/dynamicpackage_194_work/dynamicpackage_recovered_work')
OUT = ROOT / 'analysis' / 'documentation' / 'DynamicPackage_单星仿真_基础原理与详细设计_H0-H20.docx'
CJK = dds.CJK
CJK_SERIF = dds.CJK_SERIF
BLUE = dds.BLUE
TEAL = dds.TEAL


def font(run, size=10, bold=False, color=None, name=CJK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text, size=10.3, before=0, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.23
    r = p.add_run(text)
    font(r, size=size)
    return p


def principle(doc, no, name, principle_text, variables, assumptions, map_to, formulas, code_lines=None):
    dds.base.add_heading(doc, f'P{no}. {name}', 1)
    paragraph(doc, principle_text)
    dds.base.add_table(doc, ['原理维度', '说明'], [
        ('物理/数学对象', variables),
        ('基本假设与适用条件', assumptions),
        ('在当前模型中的离散实现', map_to),
    ], widths=[3.6, 13.6], font_size=8.5)
    for key, latex, caption in formulas:
        dds.base.add_formula_image(doc, key, latex, caption)
    if code_lines:
        dds.add_code(doc, f'原理到实现：{name}', code_lines)


def add_cover(doc):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(57); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('DynamicPackage'); font(r, name=CJK_SERIF, size=28, bold=True, color=BLUE)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('单星仿真基础原理与详细设计说明书'); font(r, name=CJK_SERIF, size=21, bold=True, color=TEAL)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('从物理/数学基础公式到当前 C11 行为等价模型实现｜H0–H20'); font(r, size=10.5, color='666666')
    doc.add_paragraph()
    dds.base.add_callout(doc, '文档定位', '本书在详细设计之前增加“基础原理”部分。每一个单元先说明其物理或数学起点、符号、单位、适用条件和连续形式，再说明当前 DynamicPackage 如何离散化、如何放入状态/对象、以及为何必须保留实现顺序。后半部分保留逐模块输入输出、对象布局和伪代码。', fill='E2F0D9')
    dds.base.add_table(doc, ['阅读层次', '回答的问题'], [
        ('基础原理', '为什么会有这条方程？每个变量代表什么？哪些近似成立？'),
        ('当前模型映射', '该原理在 33 维状态、设备对象、环境函数和单步调度的哪个位置实现？'),
        ('详细设计', '调用什么函数、读写什么对象、按什么顺序更新、怎样打包输出？'),
        ('行为验证', '哪些输入输出已经与原 ELF 逐步 bitwise 对照，哪些仍是有限覆盖？'),
    ], widths=[3.7, 13.5], font_size=9.0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
    r = p.add_run('仅限离线研究与仿真；不构成飞行、控制、实体执行机构或安全关键用途的设计保证。'); font(r, size=9.5, bold=True, color='9C0006')
    doc.add_page_break()


def add_principles(doc):
    dds.base.add_heading(doc, '第一部分：单星仿真基础原理', 1)
    paragraph(doc, '这一部分使用标准航天动力学、刚体动力学、数值分析和测量模型语言解释当前工程。标准原理用于理解公式来源；DynamicPackage 的具体常量、符号方向、分支、数据布局、运算顺序和可验证行为仍以工程源码与原 ELF gold 为准。', after=7)
    dds.base.add_table(doc, ['单元', '基础原理', '详细设计对应章节'], [
        ('P1', '连续状态到离散步进', '第 1、6、7、23 章'),
        ('P2', '向量、坐标、旋转与四元数', '第 2、3、5、14–17 章'),
        ('P3', '时间、天文时间与参考系', '第 3、4、5、20 章'),
        ('P4', '轨道两体问题、开普勒和带谐摄动', '第 4 章'),
        ('P5', '刚体角动量、外力矩与姿态', '第 6 章'),
        ('P6', '柔性模态与刚柔耦合', '第 6 章'),
        ('P7', 'RK4 与数值误差/归一化', '第 7 章'),
        ('P8', '执行机构的动量、磁矩、推力和限幅', '第 8–12 章'),
        ('P9', '传感器测量、噪声、可见性与有效标志', '第 13–19 章'),
        ('P10', '二进制遥测、序列化与差分验证', '第 20–23 章'),
    ], widths=[2.0, 8.2, 7.0], font_size=8.3)

    principle(doc, '1', '连续动力学为何需要离散单步模型',
              '卫星状态在物理上由连续时间微分方程描述，但软件只能在离散时刻计算。因此模型把 t、状态 y(t)、命令 u(t)、环境 e(t) 和设备状态 xdev(t) 按步长 Δt 采样。离散传播器的任务是在假设本步内命令/配置按当前调用规则生效的前提下，近似连续流映射 y(tk) 到 y(tk+1)。',
              '状态 y∈R^33；时间 t；步长 Δt；外部命令 u；设备对象 xdev；环境 e；测量 z；输出 o。',
              '连续动力学在单步内足够平滑；当前实现使用固定步长 RK4，而非事件驱动积分或自适应步长。外部命令由调用者给定，不由模型自主生成。',
              '当前模型严格执行 D(设备) → ΦRK4(33维传播) → H(传感器) → G(遥测/IPC)，设备更新先于积分，测量/输出后于积分。',
              [('p_continuous', r'\dot{\mathbf y}(t)=\mathbf f(t,\mathbf y(t),\mathbf u(t),\mathbf e(t),\Theta),\qquad \mathbf y_{k+1}\approx\Phi_{\Delta t}(\mathbf y_k,\mathbf u_k,\mathbf e_k,\Theta)', '式 P1-1：连续状态方程及其离散流映射。'),
               ('p_output', r'\mathbf x_{dev,k+1}=D(\mathbf x_{dev,k},\mathbf u_k),\quad \mathbf z_{k+1}=H(\mathbf y_{k+1},\mathbf e_{k+1},\boldsymbol\xi_k),\quad \mathbf o_{k+1}=G(\mathbf y_{k+1},\mathbf x_{dev,k+1},\mathbf z_{k+1})', '式 P1-2：当前模型的设备—状态—测量—输出分层。')],
              ['read command frame u[k]', 'update devices xdev[k] -> xdev[k+1]', 'integrate y[k] -> y[k+1] with the updated device terms', 'derive environment/measurements z[k+1] and serialize o[k+1]'])

    principle(doc, '2', '向量、坐标变换、旋转矩阵与四元数',
              '空间位置、速度、角速度、力、力矩、磁场和太阳方向都是三维向量。坐标变换的本质是同一几何向量在不同基底下的分量转换；纯旋转矩阵保持长度与夹角。四元数用四个参数描述三维旋转，避免欧拉角局部奇异，同时必须保持单位范数。',
              '向量 a,b∈R^3；方向余弦矩阵 A∈R^(3×3)；四元数 q=[q0,q1,q2,q3]T；本体系角速度 ω。',
              '旋转矩阵应正交且 det(A)=+1；四元数理论上满足 ||q||=1。当前代码用被动旋转、标量在前的四元数约定。',
              'q 位于 y[0..3]，ω 位于 y[4..6]；quat_xi/quat_psi 生成矩阵，积分后按当前顺序重新归一化。传感器安装矩阵用于把本体系真值投影到各设备坐标。',
              [('p_cross_dot', r'\mathbf a\cdot\mathbf b=\sum_{i=1}^{3}a_i b_i,\qquad \mathbf a\times\mathbf b=(a_yb_z-a_zb_y,\;a_zb_x-a_xb_z,\;a_xb_y-a_yb_x)^T', '式 P2-1：点积给出投影/夹角；叉积给出垂直方向和力矩。'),
               ('p_rotation', r'\mathbf A^T\mathbf A=\mathbf I,\qquad \mathbf v_B=\mathbf A_{B\leftarrow I}\mathbf v_I,\qquad \mathbf v_I=\mathbf A_{B\leftarrow I}^T\mathbf v_B', '式 P2-2：正交旋转的正逆变换关系。'),
               ('p_quat', r'\dot{\mathbf q}=\frac{1}{2}\Xi(\mathbf q)\boldsymbol\omega,\qquad \|\mathbf q\|=\sqrt{q_0^2+q_1^2+q_2^2+q_3^2},\qquad \mathbf q\leftarrow\mathbf q/\|\mathbf q\|', '式 P2-3：单位四元数运动学和数值归一化。')])

    principle(doc, '3', '时间、儒略日、恒星时与惯性/地固参考系',
              '轨道和环境模型同时需要“日历时间”和“连续角度时间”。儒略日 JD 把日历映射为连续天数；GMST 表示地球相对惯性空间的自转角。GCI 是近似惯性坐标系，ECEF 随地球转动。地球重力带谐和局地地磁适合在 ECEF/局地计算，轨道状态则以 GCI 表示。',
              '年 y、月 m、日 d、时 h、分 min、秒 s；JD；儒略世纪 T；GMST 角 θ；GCI/ECEF 向量。',
              '当前模型使用固定的历法/JD 和岁差多项式；TimeAdd 的秒残量和 mktime 归一化策略属于模型行为。',
              'dynamic_time.c 的 Calc_JD、Calc_theta_GMST、P(T)、Rz(θ) 构造转换；跨年与闰日已在 H12/H15/H16 对照。',
              [('p_jd', r'JD=367y-\lfloor1.75(y+\lfloor(m+9)/12\rfloor)\rfloor+\lfloor275m/9\rfloor+d+1721013.5+\frac{h+(min+s/60)/60}{24}', '式 P3-1：当前日历到儒略日的计算。'),
               ('p_gmst', r'\theta_{GMST}=\mathrm{fmod}(24110.54841+8640184.812866T+0.093104T^2-0.0000062T^3+1.0027379093507951t_{UTC},86400)\frac{\pi}{43200}', '式 P3-2：当前 GMST 多项式及弧度转换。'),
               ('p_frames', r'\mathbf r_{ECEF}=R_Z(\theta_{GMST})\mathbf P(T)\mathbf r_{GCI}', '式 P3-3：当前 GCI 到 ECEF 的旋转链。')])

    principle(doc, '4', '轨道力学：牛顿引力、开普勒椭圆与带谐摄动',
              '二体问题来自万有引力与牛顿第二定律：卫星加速度指向地心且大小与距离平方成反比。椭圆轨道可用六要素表示；偏近点角 E 是在椭圆参数化中把时间（平均近点角 M）与位置连接起来的中间变量。J2/J3/J4 是地球非球形重力势的高阶修正。',
              'r、v、a；μ；半长轴 a；偏心率 e；倾角 i；RAAN Ω；近地点幅角 ω；M/E/ν；Re/J2/J3/J4。',
              '该模型实现椭圆要素路径；近圆/近赤道有显式退化分支。带谐项在 ECEF 计算后旋回 GCI，外力按 F/m 叠加。',
              'orbit_dynamic 输出 y[10..12] 的导数；Elements2PosVel_M 用于初始化；PosVel2Elements 用于主遥测。',
              [('p_newton', r'\mathbf F=-\frac{\mu m}{r^3}\mathbf r,\qquad \ddot{\mathbf r}=\mathbf a_0=-\frac{\mu}{r^3}\mathbf r,\qquad \dot{\mathbf r}=\mathbf v', '式 P4-1：二体引力与位置/速度状态方程。'),
               ('p_energy', r'\varepsilon=\frac{\|\mathbf v\|^2}{2}-\frac{\mu}{r}=-\frac{\mu}{2a},\qquad n=\sqrt{\frac{\mu}{a^3}},\qquad P=2\pi\sqrt{\frac{a^3}{\mu}}', '式 P4-2：轨道比机械能、平均角速度和椭圆周期。'),
               ('p_kepler', r'M=E-e\sin E,\qquad E_{n+1}=E_n+\frac{M-(E_n-e\sin E_n)}{1-e\cos E_n}', '式 P4-3：开普勒方程与当前 Newton 迭代。'),
               ('p_harmonic', r'\mathbf a=\mathbf a_0+\mathbf a_{J2}+\mathbf a_{J3}+\mathbf a_{J4}+\mathbf F_{external}/m', '式 P4-4：当前轨道加速度的力学合成。')])

    principle(doc, '5', '刚体姿态动力学、角动量与外力矩',
              '转动版牛顿定律是角动量定理：惯性系角动量变化率等于外力矩。把它改写到随本体转动的坐标系，就出现 ω×H 项。飞轮通过改变内部角动量对本体产生反作用力矩；磁力矩器的偶极矩与地磁叉乘；重力梯度由惯量椭球在非均匀引力场中的取向差异产生。',
              'J、ω、α、H、τ；飞轮角动量 Hw；磁矩 m；磁场 B；本体系径向单位向量 rhatB。',
              'J 在本体系下为当前配置；无外力矩时总角动量守恒。当前模型使用已恢复的外力矩项和设备耦合项，不自行添加未取证的控制力矩。',
              'dp_differential_equation_33/torque helpers 形成 b 与 α；主桥接在积分前注入 Hwheel、minus torque 与 prior B。',
              [('p_euler', r'\frac{d\mathbf H}{dt}\bigg|_I=\boldsymbol\tau_{ext},\qquad \mathbf H=\mathbf J\boldsymbol\omega+\mathbf H_{add},\qquad \mathbf J\dot{\boldsymbol\omega}=\boldsymbol\tau_{ext}-\boldsymbol\omega\times\mathbf H', '式 P5-1：本体系欧拉方程和附加角动量。'),
               ('p_wheel_torque', r'\boldsymbol\tau_{RW}=-\dot{\mathbf H}_w', '式 P5-2：飞轮加速对本体的反作用力矩。'),
               ('p_mag_torque', r'\boldsymbol\tau_{MTQ}=\mathbf m\times\mathbf B', '式 P5-3：磁偶极与地磁的力矩。'),
               ('p_gravity_gradient', r'\boldsymbol\tau_{gg}=\frac{3\mu}{r^3}\left[\hat{\mathbf r}_B\times(\mathbf J\hat{\mathbf r}_B)\right]', '式 P5-4：重力梯度力矩。')])

    principle(doc, '6', '柔性模态、阻尼、刚柔耦合与有效惯量',
              '连续柔性结构可由无限多个变形自由度描述。模态截断把位移场展开为有限个模态形状与广义坐标 η，使大系统变为有限维常微分方程。刚体角加速度会激励模态，模态反过来改变等效惯量和角动量；这就是刚柔耦合。',
              '模态坐标 η、ηdot、ηddot∈R^10；质量 M、阻尼 D、刚度 K 的等效矩阵；耦合矩阵 C∈R^(3×10)；Jeff。',
              '当前模型固定截断为 10 阶，而不是通用有限元求解；MA/MD、C 与 SADA 映射由默认配置给出。其物理单位以模型配置为准。',
              'η 位于 y[13..22]，ηdot 位于 y[23..32]；dynamic_flex.c 先构造 Jeff，再形成刚体 RHS、alpha 和模态 RHS。',
              [('p_modal', r'\mathbf M\ddot{\boldsymbol\eta}+\mathbf D\dot{\boldsymbol\eta}+\mathbf K\boldsymbol\eta=\mathbf Q_{rigid}+\mathbf Q_{SADA}', '式 P6-1：一般模态结构动力学；当前 MA/MD 表示其恢复配置中的刚度/阻尼相关项。'),
               ('p_effective_inertia', r'\mathbf J_{eff}=\mathbf J-\mathbf C\mathbf C^T,\qquad \boldsymbol\alpha=\mathbf J_{eff}^{-1}\mathbf b', '式 P6-2：当前模型用于耦合求解的有效惯量。'),
               ('p_modal_current', r'\ddot{\boldsymbol\eta}=-\mathbf C^T\boldsymbol\alpha-0.1\mathbf M_A\boldsymbol\eta-\mathbf M_D\dot{\boldsymbol\eta}+\mathrm{optional\;SADA/rigid\;terms}', '式 P6-3：当前模态加速度的主项。')])

    principle(doc, '7', '数值积分：RK4、局部截断误差与四元数约束',
              '连续 ODE 通常没有闭式解，RK4 通过在一个步长内四次采样斜率构造高阶加权平均。对于光滑 RHS，经典 RK4 的单步局部截断误差为 O(h^5)，全局误差为 O(h^4)。数值积分会使四元数偏离单位球，因此当前模型每步归一化。',
              'f(t,y)、步长 h、四个斜率 k1..k4、q 范数。',
              '误差阶结论要求 RHS 足够光滑且步长充分小；当前恢复不使用自适应步长或误差控制。为与 ELF 对齐，浮点运算顺序比理论等价改写更重要。',
              'dynamic_rk4.c 将 h·ki 缓冲，最终按 k4→2k3→2k2→k1 顺序累加，随后 q/=||q||。',
              [('p_rk4_principle', r'\mathbf k_1=\mathbf f(t,\mathbf y),\;\mathbf k_2=\mathbf f(t+h/2,\mathbf y+h\mathbf k_1/2),\;\mathbf k_3=\mathbf f(t+h/2,\mathbf y+h\mathbf k_2/2),\;\mathbf k_4=\mathbf f(t+h,\mathbf y+h\mathbf k_3)', '式 P7-1：RK4 的四个斜率采样。'),
               ('p_rk4_update', r'\mathbf y^+=\mathbf y+\frac{h}{6}(\mathbf k_1+2\mathbf k_2+2\mathbf k_3+\mathbf k_4),\qquad \mathbf q^+=\mathbf q^+/\|\mathbf q^+\|', '式 P7-2：RK4 更新与四元数约束投影。')])

    principle(doc, '8', '执行机构：动量交换、磁偶极、推力和速率/加速度限幅',
              '飞轮通过内部转子角动量交换实现姿态力矩；MTQ 通过磁偶极与地磁交互；推力器直接施加力并经力臂产生力矩；SADA 是带位置/速率/加速度约束的双轴驱动。它们都具有离散对象状态，因此命令值不等于立即无限制的物理输出。',
              '轮惯量 Jw、轮速 ωw、轮扭矩 τw；磁矩 m；推力 F；力臂 l；SADA 角 θ、角速度 ωS、角加速度 aS。',
              '命令须经过当前对象给定的扭矩、速度、磁矩、状态和加速度限制；H18–H20 已验证多执行机构的终态，但未穷尽所有限制组合。',
              '第 8–12 章分别实现设备对象；dp_update_device_control 规定其顺序为惯量→轮→MTQ→推力器→SADA。',
              [('p_actuator_wheel', r'\tau_w=J_w\alpha_w,\qquad \omega_w^+=\omega_w+\alpha_w\Delta t,\qquad H_w^+=J_w\omega_w^+', '式 P8-1：飞轮转子离散动量更新。'),
               ('p_actuator_thruster', r'\mathbf F_{out}=k_F\mathbf F_{in},\qquad \boldsymbol\tau_{thr}=\mathbf l\times\mathbf F_{out}', '式 P8-2：推力器力和力臂力矩。'),
               ('p_actuator_sada', r'\Delta\omega_S=\mathrm{sat}(\theta_{cmd}-\omega_S,[-a_{max}\Delta t,a_{max}\Delta t]),\quad \omega_S^+=\omega_S+\Delta\omega_S,\quad \theta^+=\theta+\omega_S^+\Delta t', '式 P8-3：当前 SADA 的增量限幅离散更新。')])

    principle(doc, '9', '传感器：状态到测量、随机误差与可见性几何',
              '传感器模型的通用形式是把真实状态经安装/投影关系 h(x) 映射到测量，再叠加偏置与随机噪声。陀螺测角速度、磁强计测磁场、DSS 测太阳方向、STS 测姿态、GPS 接口复制传播的位置/速度/时间。有效标志并非噪声，而是由视场、太阳排除、地影和地球遮挡等几何条件决定。',
              '真值 x、投影矩阵 P、测量 z、偏置 b、噪声 v、太阳 s、位置 r、阈值/安装矩阵。',
              '随机噪声在当前平台依赖 C rand()；高斯模型只说明统计意图，不保证任意长样本的分布性质。DSS/STS 几何公式假定位置不在地球内部且向量范数有效。',
              '第 13–19 章实现高斯、两路 Gyro/Mag/DSS、三路 STS、GPS 和总装；DSS 有效性使用测量与 (0,0,1) 的夹角，已由 H11 取证。',
              [('p_measurement', r'\mathbf z=\mathbf h(\mathbf x)+\mathbf b+\mathbf v,\qquad \mathbf v\sim\mathcal N(\mathbf0,\mathbf R)', '式 P9-1：一般非线性测量与加性噪声模型。'),
               ('p_gaussian', r'x,y\sim(-1,1),\;s=x^2+y^2\in(0,1],\qquad n=y\sigma\sqrt{\frac{-2\ln s}{s}},\qquad z=\mu+n', '式 P9-2：当前极坐标拒绝采样高斯生成。'),
               ('p_dss', r'\mathbf u=\mathbf P_{DSS}\mathbf s_B,\;u_z>0:\quad \alpha_x=\mathrm{atan2}(u_x/u_z,1),\quad \alpha_y=\mathrm{atan2}(u_y/u_z,1)', '式 P9-3：DSS 投影和双角测量。'),
               ('p_visibility', r'\alpha_{limb}=\arccos(R_e/\|\mathbf r\|),\qquad shadow=[\angle(\mathbf s_{GCI},\mathbf r_{GCI})>\alpha_{limb}+\pi/2]', '式 P9-4：当前地影几何判定。')])

    principle(doc, '10', '二进制遥测、共享内存与行为等价验证',
              '仿真内部的双精度物理状态必须被映射为外部协议字节。序列化不仅包含数值，也包含字段偏移、float 降精度、int/u8 标志和端序。行为等价验证将原 ELF 视为黑盒金标，在相同输入轨迹上逐步比较这些字节流，因此把“模型正确性”转为可审计的实验命题。',
              'DpMainTelemetryFrame 544 B；IPC shared 0xBF0 B；偏移/端序标志；原 ELF/恢复端 gold。',
              '协议结果依赖当前 Linux x86-64 的 pthread rwlock 布局和 C 浮点/随机数；H0–H20 仅是有限场景而非全状态空间证明。',
              '第 20–23 章使用 put/memcpy/float/u8 映射与 make selftest；双采集保障 gold 自身确定性。',
              [('p_serialization', r'\mathrm{bytes}[off:off+n]=\mathrm{encode}_{endian,type}(value),\qquad value_{float}=\mathrm{float32}(value_{double})', '式 P10-1：二进制协议编码的抽象表达；实际偏移和类型以 current layout 为准。'),
               ('p_difference', r'\Delta_{k,b}=\mathrm{memcmp}(gold_{k,b},recovered_{k,b}),\qquad PASS\Leftrightarrow\forall k,b:\Delta_{k,b}=0', '式 P10-2：逐步逐块 bitwise 差分的验收定义。')])

    dds.base.add_heading(doc, '基础原理参考与实现边界', 1)
    paragraph(doc, '标准原理的背景参考可见 NASA NTRS 的《Spacecraft Modeling, Attitude Determination, and Control》；它覆盖轨道动力学、旋转、四元数和参考系等通用概念。[R1] 当前工程的模型选择、常量、函数调用顺序、数据布局和验收结论以本项目内源码/反汇编/原 ELF gold 为唯一行为证据。')
    dds.base.add_table(doc, ['编号', '来源/用途'], [
        ('[R1]', 'NASA Technical Reports Server：Spacecraft Modeling, Attitude Determination, and Control，https://ntrs.nasa.gov/citations/20240009554。用于通用轨道、姿态、旋转和参考系背景。'),
        ('[S1]–[S18]', '工程内动态数学、时间、轨道、环境、设备、传感器、动力学、RK4、遥测、IPC 与 H0–H20 取证文件。具体入口见后半部分 DDS 的源码—章节追踪矩阵。'),
    ], widths=[2.0, 15.2], font_size=8.2)
    dds.base.add_callout(doc, '必须保持的区分', '“基础公式”说明物理或数学原理；“当前实现映射”说明本恢复模型实际执行的常量、分支、数据布局和浮点顺序。二者通常相关，但不能用教材公式替代对原 ELF 行为的差分取证。', fill='FCE4D6')
    doc.add_page_break()


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.55); section.right_margin = Cm(1.55)
    dds.base.create_styles(doc)
    footer = section.footer.paragraphs[0]; footer.text = ''
    dds.base.add_page_number(footer)
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run('DynamicPackage 单星仿真基础原理与详细设计  |  H0–H20'); font(r, size=8, color='7F7F7F')

    add_cover(doc)
    add_principles(doc)
    dds.add_reader_contract(doc)
    dds.add_architecture(doc)
    dds.add_math_module(doc)
    dds.add_time_coordinate_module(doc)
    dds.add_orbit_environment_modules(doc)
    dds.add_dynamics_modules(doc)
    dds.add_device_modules(doc)
    dds.add_sensor_modules(doc)
    dds.add_output_modules(doc)
    dds.add_verification(doc)
    dds.add_appendices(doc)

    doc.core_properties.title = 'DynamicPackage 单星仿真基础原理与详细设计说明书（H0–H20）'
    doc.core_properties.subject = '基础公式、原理、当前模型映射与模块化详细设计'
    doc.core_properties.author = 'Manus AI'
    doc.core_properties.comments = '仅限离线研究与仿真。'
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build_doc()
